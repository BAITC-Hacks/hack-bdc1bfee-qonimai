"""Pure UI state transitions and local export, without Streamlit or model calls."""

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zoneinfo import ZoneInfo

from contracts import validate_analysis, validate_export, validate_snapshot, validate_transcript


class UIValidationError(ValueError):
    """Fixed, user-safe errors; never include private input or paths."""


def retain_snapshot(history, transcript, analysis=None, *, pending_analysis=None,
                    reviewed_revision=None, synthetic=False, synthetic_audio=False):
    """Preserve the complete previous state before replacing any input/result.

    History is session memory only. Never truncate it silently; the UI exposes
    JSON download and a separate explicit deletion action.
    """
    result = deepcopy(history)
    if transcript is None:
        return result
    snapshot = {"transcript": deepcopy(transcript), "analysis": deepcopy(analysis),
                "pending_analysis": deepcopy(pending_analysis),
                "reviewed_revision": reviewed_revision,
                "synthetic": bool(synthetic), "synthetic_audio": bool(synthetic_audio)}
    if not result or result[-1] != snapshot:
        result.append(snapshot)
    return result


def participants_from_text(text):
    names = [line.strip() for line in text.splitlines() if line.strip()]
    if len(names) > 30 or any(len(name) > 100 for name in names):
        raise UIValidationError("Не более 30 участников, имя не длиннее 100 символов.")
    return [{"id": f"p{i + 1}", "name": name} for i, name in enumerate(names)]


def meeting_metadata(day, clock, timezone, known=True):
    if not known:
        return {"started_at": None, "timezone": None}
    try:
        zone = ZoneInfo(timezone.strip())
        started = datetime.combine(day, clock).replace(tzinfo=zone)
        # Nonexistent local wall times must not be silently normalised.
        from datetime import timezone as datetime_timezone
        restored = started.astimezone(datetime_timezone.utc).astimezone(zone)
        if restored.replace(tzinfo=None) != started.replace(tzinfo=None):
            raise ValueError()
        if started.replace(fold=0).utcoffset() != started.replace(fold=1).utcoffset():
            raise ValueError()
        return {"started_at": started.isoformat(), "timezone": timezone.strip()}
    except Exception:
        raise UIValidationError("Проверьте дату, время и IANA-зону; неоднозначное местное время не поддерживается.") from None


def replace_transcript(transcript, speaker_mapping, segment_texts):
    if validate_transcript(transcript):
        raise UIValidationError("Исходный транскрипт не прошёл проверку.")
    result = deepcopy(transcript)
    people = {p["id"]: p["name"] for p in result["participants"]}
    if set(speaker_mapping) != {s["speaker_id"] for s in result["speakers"]}:
        raise UIValidationError("Таблица говорящих неполна.")
    if set(segment_texts) != {s["id"] for s in result["segments"]}:
        raise UIValidationError("Набор реплик изменён; редактируйте только текст.")
    for speaker in result["speakers"]:
        person = speaker_mapping[speaker["speaker_id"]]
        if person is not None and person not in people:
            raise UIValidationError("Используйте участника из подтверждённого списка.")
        speaker.update(participant_id=person, confirmed=person is not None,
                       display_name=people.get(person))
    for segment in result["segments"]:
        value = segment_texts[segment["id"]]
        if not isinstance(value, str) or not value.strip():
            raise UIValidationError("Текст реплики не должен быть пустым.")
        segment["text"] = value.strip()
    if result != transcript:
        result["revision"] += 1
    if validate_transcript(result):
        raise UIValidationError("Исправленный транскрипт не прошёл проверку.")
    return result


def _reasons(task, original):
    assignee_codes = {"unknown_assignee", "ambiguous_assignee", "unconfirmed_assignee"}
    due_codes = {"missing_due", "ambiguous_due", "missing_meeting_context"}
    reasons = []
    if task["assignee_id"] is None:
        prior = [r for r in original["clarification_reasons"] if r["code"] in assignee_codes]
        reasons.extend(deepcopy(prior) or [{"code": "unknown_assignee", "message": "Исполнитель требует уточнения."}])
    if task["due_date"] is None:
        prior = [r for r in original["clarification_reasons"] if r["code"] in due_codes]
        reasons.extend(deepcopy(prior) or [{"code": "missing_due" if task["due_text"] is None else "ambiguous_due",
                                          "message": "Срок требует уточнения."}])
    task["clarification_reasons"] = reasons
    task["requires_clarification"] = bool(reasons)


def apply_user_edits(transcript, analysis, changes, reason):
    """Return copies; changed snapshots are stale until content revalidation.

    changes: list of {target_type, target_id, values}. Only user-editable fields.
    This never sets reviewed_revision or silently overwrites the first original.
    """
    if validate_analysis(transcript, analysis) or analysis["status"] != "ok":
        raise UIValidationError("Результат анализа не допускает редактирование.")
    if not isinstance(reason, str) or not reason.strip():
        raise UIValidationError("Укажите объяснение исправлений.")
    updated = deepcopy(analysis)
    data = updated["data"]
    names = {"task": "tasks", "summary": "summary", "decision": "decisions"}
    allowed = {"task": {"description", "assignee_id", "due_date"}, "summary": {"text"}, "decision": {"text"}}
    seen = set()
    for change in changes:
        kind, identity, values = change.get("target_type"), change.get("target_id"), change.get("values")
        if kind not in names or not isinstance(identity, str) or (kind, identity) in seen:
            raise UIValidationError("Некорректная цель пользовательской правки.")
        seen.add((kind, identity))
        if not isinstance(values, dict) or not set(values) <= allowed[kind]:
            raise UIValidationError("Изменение исходных ссылок и сроков в этом редакторе запрещено.")
        current = next((row for row in data[names[kind]] if row["id"] == identity), None)
        if current is None:
            raise UIValidationError("Изменяемый объект не найден.")
        existing = next((e for e in data["user_edits"] if e["target_type"] == kind and e["target_id"] == identity), None)
        original = deepcopy(existing["original"] if existing else current)
        current.update(deepcopy(values))
        if kind == "task":
            if current["due_date"] is not None:
                try:
                    if date.fromisoformat(current["due_date"]).isoformat() != current["due_date"]:
                        raise ValueError()
                except (TypeError, ValueError):
                    raise UIValidationError("Дата поручения должна иметь вид ГГГГ-ММ-ДД или быть пустой.") from None
            _reasons(current, original)
            if all(current[k] == original[k] for k in allowed[kind]):
                current.update(deepcopy(original))
        fields = [key for key in original if current.get(key) != original[key]]
        data["user_edits"] = [e for e in data["user_edits"] if not (e["target_type"] == kind and e["target_id"] == identity)]
        if fields:
            data["user_edits"].append({"target_type": kind, "target_id": identity, "original": original,
                                       "changed_fields": fields, "reason": reason.strip()})
    if validate_analysis(transcript, updated):
        raise UIValidationError("Исправления не прошли проверку. Проверьте поля и ссылки на источники.")
    new_transcript = deepcopy(transcript)
    if updated != analysis:
        new_transcript["revision"] += 1
    return new_transcript, updated


def accept_content_review(transcript, analysis, confirmed=False):
    """User separately reviewed sources; validate before accepting new revision."""
    if confirmed is not True:
        raise UIValidationError("Сначала сверьте изменённый результат с репликами и подтвердите проверку.")
    if validate_analysis(transcript, analysis):
        raise UIValidationError("Данные не прошли проверку; подтверждение не исправляет ошибки.")
    result = deepcopy(analysis)
    result["data"]["source_revision"] = transcript["revision"]
    if validate_snapshot(transcript, result):
        raise UIValidationError("Текущий снимок не прошёл проверку.")
    return result


def export_document_bytes(transcript, analysis, reviewed_revision, synthetic=False, synthetic_audio=False):
    """Generate in a disposable directory; return bytes, never retain transcript."""
    if validate_export(transcript, analysis, reviewed_revision):
        raise UIValidationError("Для экспорта нужен актуальный проверенный снимок и подтверждение просмотра.")
    from export_docx import export_docx
    with TemporaryDirectory(prefix="qonimai-export-") as directory:
        destination = Path(directory) / "protocol.docx"
        result = export_docx(transcript, analysis, reviewed_revision, destination)
        if result["status"] != "ok":
            raise UIValidationError("Не удалось создать DOCX в локальном временном каталоге.")
        content = destination.read_bytes()
    if synthetic or synthetic_audio:
        from docx import Document
        document = Document(BytesIO(content))
        notice = document.add_paragraph("Синтетический пример — модели не запускались" if synthetic
                                        else "Синтетическая аудиозапись — выполнена локальная обработка моделями")
        notice.runs[0].bold = True
        document.paragraphs[0]._p.addnext(notice._p)
        output = BytesIO()
        document.save(output)
        content = output.getvalue()
    return content
