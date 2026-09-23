"""Export an approved v1 snapshot to a local DOCX, without a model or network.

The caller owns review and the destination directory. No file is touched before
contracts.validate_export succeeds. A completed temporary file replaces the
destination atomically; failures do not overwrite the previous document.
"""

import ctypes
import os
from pathlib import Path
import re
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import contracts


_ERROR_PRIORITY = (
    "TIMEZONE_DATA_UNAVAILABLE", "INVALID_INPUT", "INVALID_MODEL_OUTPUT",
    "STALE_ANALYSIS", "NO_CONTENT", "REVIEW_REQUIRED",
)
_ERROR_MESSAGES = {
    "TIMEZONE_DATA_UNAVAILABLE": "База часовых поясов IANA недоступна; требуется настройка среды.",
    "INVALID_INPUT": "Входной транскрипт не соответствует контракту.",
    "INVALID_MODEL_OUTPUT": "Результат анализа не соответствует контракту.",
    "STALE_ANALYSIS": "Результат анализа устарел; требуется повторная проверка снимка.",
    "NO_CONTENT": "Нет речи для экспорта протокола.",
    "REVIEW_REQUIRED": "Требуется просмотр пользователем текущей версии.",
    "EXPORT_FAILED": "Не удалось сохранить документ в выбранный локальный файл DOCX.",
}
_FIELD_LABELS = {
    "description": "Описание", "assignee_id": "Исполнитель", "due_date": "Нормализованная дата",
    "requires_clarification": "Необходимость уточнения", "clarification_reasons": "Причины уточнения",
    "text": "Текст",
}
_FIELD_ORDER = ("description", "assignee_id", "due_date", "requires_clarification", "clarification_reasons", "text")


def _error(code):
    return {"status": "error", "data": None,
            "error": {"code": code, "message": _ERROR_MESSAGES[code], "retryable": False}}


def _local_docx_path(target_path):
    """Reject URLs and network/device paths; never create parent directories."""
    raw = os.fspath(target_path)
    if type(raw) is not str or not raw.strip() or "\x00" in raw:
        raise ValueError("Invalid path")
    if raw.startswith(("\\\\", "//")) or "://" in raw:
        raise ValueError("Not a local path")
    if re.match(r"^[A-Za-z][A-Za-z0-9+.-]*:", raw) and not re.match(r"^[A-Za-z]:[\\/]", raw):
        raise ValueError("Not an absolute drive path")
    path = Path(raw).absolute()
    if os.name == "nt":
        if path.is_reserved() or any(":" in part for part in path.parts[1:]):
            raise ValueError("Device names and alternate streams are unsupported")
        if ctypes.windll.kernel32.GetDriveTypeW(ctypes.c_wchar_p(path.anchor)) == 4:
            raise ValueError("Mapped network drive is unsupported")
    if path.suffix.lower() != ".docx" or path.is_symlink():
        raise ValueError("Not a regular DOCX destination")
    # Avoid resolving a symlinked parent to a different (possibly remote) place.
    if any(parent.is_symlink() or (hasattr(parent, "is_junction") and parent.is_junction()) for parent in path.parents):
        raise ValueError("Linked parent is unsupported")
    if not path.parent.is_dir() or (path.exists() and not path.is_file()):
        raise ValueError("Destination directory is unavailable")
    return path


def _word_text(value):
    # Treat CRLF and standalone CR as logical line breaks; preserve tabs and
    # every other character. Illegal XML control characters cause EXPORT_FAILED.
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _paragraph(doc, text="", style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(_word_text(text))
    return paragraph


def _label(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.add_run(label + " ").bold = True
    paragraph.add_run(_word_text(value))
    return paragraph


def _person(participant_id, people):
    if participant_id is None:
        return "не определён"
    return f"{people[participant_id]['name']} ({participant_id})"


def _speaker(speaker_id, speakers, people):
    speaker = speakers[speaker_id]
    if speaker["confirmed"]:
        return f"{_person(speaker['participant_id'], people)}; спикер {speaker_id}"
    label = f"Спикер {speaker_id} (личность не подтверждена)"
    if speaker["display_name"] is not None:
        label += f"; неподтверждённая подсказка имени: {speaker['display_name']}"
    return label


def _value(field, value, people):
    if field == "assignee_id":
        return _person(value, people)
    if field == "requires_clarification":
        return "требуется" if value else "не требуется"
    if field == "clarification_reasons":
        return "; ".join(reason["message"] for reason in value) or "причин нет"
    if value is None:
        return "не указана" if field == "due_date" else "не указано"
    return value


def _user_edit(doc, current, edit, people):
    """Display original and current values for every permitted target type."""
    if edit is None:
        return
    _paragraph(doc, "Уточнено пользователем", "Heading 3")
    for field in _FIELD_ORDER:
        if field in edit["changed_fields"]:
            _paragraph(doc, _FIELD_LABELS[field], "Heading 4")
            _label(doc, "Исходное значение:", _value(field, edit["original"][field], people))
            _label(doc, "Актуальное значение:", _value(field, current[field], people))
    _label(doc, "Объяснение пользователя:", edit["reason"])


def _sources(doc, item):
    _label(doc, "Источники:", ", ".join(item["source_segment_ids"]))


def _configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = False
    for name, size in (("Title", 22), ("Subtitle", 11), ("Heading 1", 16),
                       ("Heading 2", 13), ("Heading 3", 11), ("Heading 4", 11)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        if name.startswith("Heading"):
            style.font.bold = True
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[name]
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Arial")
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        color = style.element.rPr.find(qn("w:color"))
        if color is not None:
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + attr), None)
    # A page number helps navigate long transcripts without using decorative UI.
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run("Страница ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Протокол совещания"
    doc.core_properties.subject = "Проверенный снимок совещания"
    doc.core_properties.author = "QonimAI"
    doc.core_properties.last_modified_by = "QonimAI"
    doc.core_properties.comments = ""


def _build_document(transcript, analysis):
    doc = Document()
    _configure(doc)
    data = analysis["data"]
    people = {row["id"]: row for row in transcript["participants"]}
    speakers = {row["speaker_id"]: row for row in transcript["speakers"]}
    edits = {(row["target_type"], row["target_id"]): row for row in data["user_edits"]}
    _paragraph(doc, "Протокол совещания", "Title")
    _label(doc, "Дата совещания:", transcript["meeting"]["started_at"] or "не указана")
    _label(doc, "Часовой пояс:", transcript["meeting"]["timezone"] or "не указан")
    _paragraph(doc, "Протокол содержит саммари, решения, поручения и исходные реплики. Неизвестные значения и пользовательские уточнения отмечены отдельно.")
    _paragraph(doc, "Подтверждённые участники", "Heading 1")
    if not people:
        _paragraph(doc, "Подтверждённые участники не указаны.")
    for participant in transcript["participants"]:
        _paragraph(doc, f"{participant['name']} ({participant['id']})")
    for collection, target_type, heading, empty in (
        ("summary", "summary", "Саммари", "Саммари не указано."),
        ("decisions", "decision", "Принятые решения", "Принятые решения не зафиксированы."),
    ):
        _paragraph(doc, heading, "Heading 1")
        if not data[collection]:
            _paragraph(doc, empty)
        for item in data[collection]:
            _paragraph(doc, item["text"])
            _sources(doc, item)
            _user_edit(doc, item, edits.get((target_type, item["id"])), people)
    _paragraph(doc, "Поручения", "Heading 1")
    if not data["tasks"]:
        _paragraph(doc, "Поручения не выявлены.")
    for number, task in enumerate(data["tasks"], 1):
        _paragraph(doc, f"Поручение {number}", "Heading 2")
        _label(doc, "Идентификатор:", task["id"])
        _paragraph(doc, task["description"])
        _label(doc, "Исполнитель:", _person(task["assignee_id"], people))
        _label(doc, "Автор поручения:", _speaker(task["author_speaker_id"], speakers, people))
        _label(doc, "Исходный срок:", task["due_text"] if task["due_text"] is not None else "не указан")
        _label(doc, "Нормализованная дата:", task["due_date"] if task["due_date"] is not None else "не указана")
        if task["requires_clarification"]:
            _paragraph(doc, "Требует уточнения", "Heading 3")
            for reason in task["clarification_reasons"]:
                _paragraph(doc, reason["message"])
        else:
            _paragraph(doc, "Уточнение не требуется.")
        _sources(doc, task)
        _user_edit(doc, task, edits.get(("task", task["id"])), people)
    heading = _paragraph(doc, "Транскрипт совещания", "Heading 1")
    heading.paragraph_format.page_break_before = True
    _paragraph(doc, "Идентификаторы ниже соответствуют ссылкам на источники. Время указано в секундах от начала записи.")
    for number, segment in enumerate(transcript["segments"], 1):
        label = f"Реплика {number}  {segment['id']}  {segment['start']}–{segment['end']} с"
        _paragraph(doc, label, "Heading 2")
        _label(doc, "Говорящий:", _speaker(segment["speaker_id"], speakers, people))
        _paragraph(doc, segment["text"])
    return doc


def export_docx(transcript, analysis, reviewed_revision, target_path) -> dict:
    """Validate, write a temporary local DOCX and replace only on success."""
    errors = contracts.validate_export(transcript, analysis, reviewed_revision)
    if errors:
        codes = {error["code"] for error in errors}
        code = next((code for code in _ERROR_PRIORITY if code in codes), "INVALID_MODEL_OUTPUT")
        return _error(code)

    temporary = None
    try:
        path = _local_docx_path(target_path)
        doc = _build_document(transcript, analysis)
        # Close the handle before python-docx opens the file on Windows.
        with tempfile.NamedTemporaryFile(prefix=".qonimai-", suffix=".docx", dir=path.parent, delete=False) as stream:
            temporary = Path(stream.name)
        doc.save(str(temporary))
        # Do not publish a partial/non-DOCX payload even if a writer returned
        # without an exception. This is structural verification, not render QA.
        Document(str(temporary))
        os.replace(temporary, path)
        temporary = None
        return {"status": "ok", "data": {"path": str(path), "source_revision": analysis["data"]["source_revision"]}, "error": None}
    except Exception:
        # Format/path/write failures are private; do not expose their messages.
        # KeyboardInterrupt and SystemExit are not subclasses of Exception.
        return _error("EXPORT_FAILED")
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                # Only our own temporary path is eligible for cleanup. A locked
                # temp file may require caller cleanup; never delete the target.
                pass
