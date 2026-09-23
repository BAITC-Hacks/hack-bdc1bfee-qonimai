"""Export a reviewed fictional fixture; never load recordings or model output.

Run from the repository root with ``python -m examples.export_fixture``.
The CLI adds a synthetic notice to its own staged document, without changing
the contract snapshot or the public exporter interface.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import json
import os
from pathlib import Path
import sys
import tempfile

from docx import Document
from docx.shared import RGBColor

from export_docx import _local_docx_path, export_docx


SYNTHETIC_NOTICE = "Синтетический пример — модели не запускались"
_FIXTURE_PATH = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_cases.json"


class _PrivateArgumentParser(argparse.ArgumentParser):
    """Do not echo unknown arguments, which can contain sensitive paths."""

    def error(self, message):
        self.exit(2, "Некорректные параметры запуска. Используйте --help.\n")


def build_stress_case() -> dict:
    """Return a fresh fictional snapshot with long text and all edit targets.

    This is authored test data, not output from a model and not an alteration
    of the canonical fixture file. Kazakh letters test glyph preservation only.
    """
    glyphs = "Ә Ғ Қ Ң Ө Ұ Ү Һ І ә ғ қ ң ө ұ ү һ і"
    long_token = "СИНТЕТИЧЕСКАЯНЕПРЕРЫВНАЯСТРОКАДЛЯПРОВЕРКИПЕРЕНОСА" * 5
    discussion = (
        "В этом вымышленном обсуждении участники подробно описывают материалы: "
        "для каждого пункта нужно сохранить название, короткое объяснение и "
        "проверяемую ссылку на реплику. Неизвестные сведения остаются неизвестными. "
        "Проверяем перенос длинного текста на следующую строку и страницу, "
        "сохранение знаков «кавычки», тире — и обычных переносов строк. "
        "Этот абзац проверяет оформление документа, а не качество анализа речи."
    )
    segment_texts = [
        "Данияр, подготовь список вопросов для учебной демонстрации. " + discussion,
        "Әсел, подготовь таблицу материалов до 30 сентября 2026 года. " + discussion,
        "Нужно подготовить пояснение к макету. Исполнитель и срок пока не названы. " + discussion,
        "Әсел, проверь названия разделов к следующей пятнице. " + discussion,
        "Решили использовать один общий шаблон для учебных материалов. " + discussion,
        "Синтетическая проверка букв казахского алфавита: " + glyphs
        + "\nПеренос строки в исходной реплике должен сохраниться. " + discussion,
        "Говорящий в этом примере не подтверждён. Его подсказка имени не является установленной личностью. "
        + "\nДлинный непрерывный токен: " + long_token,
        "Я подготовлю список вопросов; дополнительных сроков в записи нет. " + discussion,
    ]
    transcript = {
        "schema_version": "qonimai.analysis.v1",
        "revision": 2,
        "meeting": {"started_at": "2026-09-23T10:00:00+05:00", "timezone": "Asia/Qyzylorda"},
        "participants": [
            {"id": "p1", "name": "Ирина Соколова"},
            {"id": "p2", "name": "Данияр Әлімов"},
            {"id": "p3", "name": "Әсел Қанатқызы"},
        ],
        "speakers": [
            {"speaker_id": "spk1", "participant_id": "p1", "display_name": "Ирина Соколова", "confirmed": True},
            {"speaker_id": "spk2", "participant_id": "p2", "display_name": "Данияр Әлімов", "confirmed": True},
            {"speaker_id": "spk3", "participant_id": None, "display_name": "Айбек?", "confirmed": False},
        ],
        "segments": [
            {
                "id": f"s{index}", "start": (index - 1) * 30, "end": index * 30,
                "speaker_id": "spk3" if index == 7 else "spk2" if index == 8 else "spk1",
                "text": content,
            }
            for index, content in enumerate(segment_texts, start=1)
        ],
    }
    tasks = [
        {
            "id": "t1", "description": "Подготовить список вопросов для учебной демонстрации",
            "author_speaker_id": "spk1", "assignee_id": "p2", "due_text": None, "due_date": None,
            "source_segment_ids": ["s1", "s8"], "requires_clarification": True,
            "clarification_reasons": [{"code": "missing_due", "message": "В исходной записи срок не назван; требуется уточнить дату."}],
        },
        {
            "id": "t2", "description": "Подготовить таблицу материалов",
            "author_speaker_id": "spk1", "assignee_id": "p3", "due_text": "до 30 сентября 2026 года", "due_date": "2026-09-30",
            "source_segment_ids": ["s2"], "requires_clarification": False, "clarification_reasons": [],
        },
        {
            "id": "t3", "description": "Подготовить пояснение к макету",
            "author_speaker_id": "spk1", "assignee_id": None, "due_text": None, "due_date": None,
            "source_segment_ids": ["s3"], "requires_clarification": True,
            "clarification_reasons": [
                {"code": "unknown_assignee", "message": "Исполнитель не назван; автора реплики нельзя автоматически назначать исполнителем."},
                {"code": "missing_due", "message": "Срок не назван; требуется уточнить дату."},
            ],
        },
        {
            "id": "t4", "description": "Проверить названия разделов",
            "author_speaker_id": "spk1", "assignee_id": "p3", "due_text": "к следующей пятнице", "due_date": None,
            "source_segment_ids": ["s4"], "requires_clarification": True,
            "clarification_reasons": [{"code": "ambiguous_due", "message": "Формулировка «к следующей пятнице» требует уточнения точной даты."}],
        },
    ]
    summary = {
        "id": "f1", "text": "Обсудили вопросы, таблицу материалов, пояснение к макету и названия разделов.",
        "source_segment_ids": ["s1", "s2", "s3", "s4"],
    }
    decision = {
        "id": "d1", "text": "Использовать один общий шаблон для учебных материалов.",
        "source_segment_ids": ["s5"],
    }
    original_task, original_summary, original_decision = deepcopy((tasks[0], summary, decision))
    tasks[0]["description"] = "Подготовить список вопросов для учебной демонстрации, сгруппировав их по темам"
    tasks[0]["due_date"] = "2026-10-02"
    tasks[0]["requires_clarification"] = False
    tasks[0]["clarification_reasons"] = []
    summary["text"] = "Обсудили подготовку учебных материалов: список вопросов, таблицу, пояснение к макету и проверку названий."
    decision["text"] = "Использовать один общий шаблон для учебных материалов с отдельным полем для пояснений."
    analysis = {
        "status": "ok",
        "data": {
            "source_revision": 2,
            "summary": [summary], "decisions": [decision], "tasks": tasks,
            "user_edits": [
                {
                    "target_type": "task", "target_id": "t1", "original": original_task,
                    "changed_fields": ["description", "due_date", "requires_clarification", "clarification_reasons"],
                    "reason": "Пользователь после просмотра уточнил группировку вопросов и отдельно назначил дату; в записи этого срока нет.",
                },
                {
                    "target_type": "summary", "target_id": "f1", "original": original_summary,
                    "changed_fields": ["text"],
                    "reason": "Пользователь уточнил формулировку саммари, сохранив предмет обсуждения и ссылки.",
                },
                {
                    "target_type": "decision", "target_id": "d1", "original": original_decision,
                    "changed_fields": ["text"],
                    "reason": "Пользователь после просмотра добавил требование к полю пояснений; исходное решение сохранено отдельно.",
                },
            ],
        },
        "error": None,
    }
    return {"id": "STRESS", "input": transcript, "expected": analysis}


def _load_case(case_id: str) -> dict | None:
    if case_id == "STRESS":
        return build_stress_case()
    fixtures = json.loads(_FIXTURE_PATH.read_text(encoding="utf-8"))
    return next((case for case in fixtures["cases"] if case["id"] == case_id), None)


def _export_synthetic(case: dict, output: str) -> bool:
    """Stage the entire labelled demo; preserve an earlier target on failure."""
    staged_path = None
    try:
        # Reuse the internal destination gate: the public exporter only sees
        # the staging path, while these same rules must protect the final path.
        target = _local_docx_path(output)
        if "synthetic" not in target.stem.lower():
            return False
        descriptor, temporary = tempfile.mkstemp(prefix=".qonimai-synthetic-", suffix=".docx", dir=target.parent)
        staged_path = Path(temporary)
        os.close(descriptor)
        result = export_docx(case["input"], case["expected"], case["input"]["revision"], staged_path)
        if result["status"] != "ok":
            return False
        document = Document(staged_path)
        # The notice is part of the delivered DOCX, as well as the CLI output
        # and filename. It does not enter the original transcript or analysis.
        notice = document.add_paragraph(SYNTHETIC_NOTICE)
        notice.runs[0].bold = True
        notice.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        notice.paragraph_format.keep_with_next = True
        document.paragraphs[0]._p.addnext(notice._p)
        document.save(staged_path)
        # A writer can return without raising after producing invalid content.
        # Reopen the completed package before publishing the final destination.
        Document(staged_path)
        os.replace(staged_path, target)
        return True
    except Exception:
        # Fixed CLI output below; never print an exception, path or input text.
        # BaseException (including interruption and exit) is not suppressed.
        return False
    finally:
        if staged_path is not None:
            try:
                staged_path.unlink(missing_ok=True)
            except OSError:
                pass


def main(argv: list[str] | None = None) -> int:
    parser = _PrivateArgumentParser(description=SYNTHETIC_NOTICE)
    parser.add_argument("--case", required=True, help="Идентификатор канонического случая или STRESS")
    parser.add_argument("--output", required=True, help="Локальный .docx; имя должно содержать synthetic; каталог уже существует")
    parser.add_argument("--confirm-reviewed", action="store_true", help="Я просмотрел выбранный синтетический снимок перед экспортом")
    arguments = parser.parse_args(argv)
    print(SYNTHETIC_NOTICE)
    if not arguments.confirm_reviewed:
        print("Экспорт отменён: после просмотра примера укажите --confirm-reviewed.", file=sys.stderr)
        return 2
    try:
        case = _load_case(arguments.case)
    except Exception:
        print("Не удалось прочитать канонические синтетические данные.", file=sys.stderr)
        return 1
    if case is None:
        print("Неизвестный синтетический случай.", file=sys.stderr)
        return 2
    if not _export_synthetic(case, arguments.output):
        print("Экспорт синтетического примера не выполнен: проверьте снимок, локальный путь .docx с synthetic в имени и существующий каталог.", file=sys.stderr)
        return 1
    print("Синтетический DOCX сохранён. Визуальная проверка страниц выполняется отдельно.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
