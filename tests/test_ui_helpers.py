import copy
from datetime import date, time
from io import BytesIO
import json
import importlib.util
from pathlib import Path
import unittest
from unittest.mock import patch

from contracts import validate_export, validate_snapshot
from ui_helpers import (UIValidationError, accept_content_review, apply_user_edits,
                        export_document_bytes, meeting_metadata, participants_from_text,
                        replace_transcript, retain_snapshot)


CASES = json.loads((Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_cases.json").read_text(encoding="utf-8"))["cases"]


def case(identity="NO_DUE"):
    value = next(c for c in CASES if c["id"] == identity)
    return copy.deepcopy(value["input"]), copy.deepcopy(value["expected"])


def edit_due(value):
    return [{"target_type": "task", "target_id": "t1", "values": {"due_date": value}}]


class UIHelperTests(unittest.TestCase):
    def test_retained_snapshot_preserves_edits_without_shared_mutable_state(self):
        transcript, analysis = case("USER_CLARIFIED_DUE")
        previous = copy.deepcopy((transcript, analysis))
        history = retain_snapshot([], transcript, analysis, reviewed_revision=transcript["revision"])
        self.assertEqual(history[0]["analysis"]["data"]["user_edits"], analysis["data"]["user_edits"])
        transcript["segments"][0]["text"] = "Другой синтетический текст"
        analysis["data"]["user_edits"].clear()
        self.assertEqual(history[0]["transcript"], previous[0])
        self.assertEqual(history[0]["analysis"], previous[1])
        self.assertEqual(retain_snapshot(history, None), history)

    def test_retained_snapshots_are_not_silently_pruned(self):
        transcript, analysis = case("USER_CLARIFIED_DUE")
        history = []
        for revision in range(1, 6):
            transcript["revision"] = revision
            history = retain_snapshot(history, transcript, analysis)
        self.assertEqual(len(history), 5)
        self.assertEqual(history[0]["transcript"]["revision"], 1)

    def test_participants_do_not_deduplicate_names_or_guess_identity(self):
        self.assertEqual(participants_from_text(" Алия\n\nАлия "), [{"id": "p1", "name": "Алия"}, {"id": "p2", "name": "Алия"}])

    def test_unknown_meeting_does_not_use_system_date(self):
        self.assertEqual(meeting_metadata(None, None, "invalid", False), {"started_at": None, "timezone": None})

    def test_meeting_timezone_and_offset(self):
        value = meeting_metadata(date(2026, 9, 23), time(10), "Asia/Qyzylorda")
        self.assertEqual(value["started_at"], "2026-09-23T10:00:00+05:00")

    def test_invalid_zone_does_not_leak_input(self):
        with self.assertRaises(UIValidationError) as error:
            meeting_metadata(date(2026, 9, 23), time(10), "secret/meeting")
        self.assertNotIn("secret", str(error.exception))

    def test_dst_ambiguous_local_time_rejected(self):
        with self.assertRaises(UIValidationError):
            meeting_metadata(date(2026, 11, 1), time(1, 30), "America/New_York")

    def test_edit_bumps_revision_but_does_not_approve_or_forge_source(self):
        transcript, analysis = case()
        before = copy.deepcopy((transcript, analysis))
        changed_t, changed_a = apply_user_edits(transcript, analysis, edit_due("2026-10-01"), "Уточнение после встречи")
        self.assertEqual((transcript, analysis), before)
        self.assertEqual(changed_t["revision"], transcript["revision"] + 1)
        self.assertEqual(changed_a["data"]["source_revision"], analysis["data"]["source_revision"])
        self.assertIn("STALE_ANALYSIS", [e["code"] for e in validate_export(changed_t, changed_a, changed_t["revision"])])
        self.assertIsNone(changed_a["data"]["tasks"][0]["due_text"])

    def test_content_review_and_export_are_separate(self):
        transcript, analysis = case()
        t, a = apply_user_edits(transcript, analysis, edit_due("2026-10-01"), "Уточнение")
        with self.assertRaises(UIValidationError):
            accept_content_review(t, a)
        approved = accept_content_review(t, a, True)
        self.assertEqual(validate_snapshot(t, approved), [])
        self.assertIn("REVIEW_REQUIRED", [e["code"] for e in validate_export(t, approved, None)])
        self.assertEqual(validate_export(t, approved, t["revision"]), [])

    def test_original_preserved_on_second_edit(self):
        transcript, analysis = case()
        t, a = apply_user_edits(transcript, analysis, edit_due("2026-10-01"), "Первая правка")
        t2, a2 = apply_user_edits(t, a, edit_due("2026-10-02"), "Вторая правка")
        self.assertEqual(a2["data"]["user_edits"][0]["original"], analysis["data"]["tasks"][0])
        self.assertEqual(len(a2["data"]["user_edits"]), 1)
        self.assertEqual(t2["revision"], t["revision"] + 1)

    def test_revert_removes_edit_and_restores_exact_reasons(self):
        transcript, analysis = case()
        t, a = apply_user_edits(transcript, analysis, edit_due("2026-10-01"), "Правка")
        _, restored = apply_user_edits(t, a, edit_due(None), "Отмена правки")
        self.assertEqual(restored["data"]["user_edits"], [])
        self.assertEqual(restored["data"]["tasks"], analysis["data"]["tasks"])

    def test_forbidden_source_edit_rejected(self):
        transcript, analysis = case()
        with self.assertRaises(UIValidationError):
            apply_user_edits(transcript, analysis, [{"target_type": "task", "target_id": "t1", "values": {"due_text": "вымысел"}}], "Правка")

    def test_invalid_date_rejected(self):
        transcript, analysis = case()
        with self.assertRaises(UIValidationError):
            apply_user_edits(transcript, analysis, edit_due("2026-02-30"), "Правка")

    def test_reason_required(self):
        transcript, analysis = case()
        with self.assertRaises(UIValidationError):
            apply_user_edits(transcript, analysis, edit_due("2026-10-01"), " ")

    def test_transcript_mapping_change_invalidates_old_analysis(self):
        transcript, analysis = case("RU_EXPLICIT")
        original = copy.deepcopy(transcript)
        mapping = {s["speaker_id"]: None for s in transcript["speakers"]}
        revised = replace_transcript(transcript, mapping, {s["id"]: s["text"] for s in transcript["segments"]})
        self.assertEqual(transcript, original)
        self.assertEqual(revised["revision"], transcript["revision"] + 1)
        self.assertTrue(all(s["participant_id"] is None and not s["confirmed"] for s in revised["speakers"]))
        self.assertIn("STALE_ANALYSIS", [e["code"] for e in validate_snapshot(revised, analysis)])

    def test_missing_segments_rejected(self):
        transcript, _ = case()
        with self.assertRaises(UIValidationError):
            replace_transcript(transcript, {s["speaker_id"]: s["participant_id"] for s in transcript["speakers"]}, {})

    def test_synthetic_export_is_marked_in_real_docx(self):
        from docx import Document
        transcript, analysis = case("USER_CLARIFIED_DUE")
        data = export_document_bytes(transcript, analysis, transcript["revision"], synthetic=True)
        text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
        self.assertIn("Синтетический пример — модели не запускались", text)
        self.assertIn("Уточнено пользователем", text)
        self.assertIn(transcript["segments"][0]["text"], text)

    def test_export_requires_review_and_rejects_stale_snapshot(self):
        transcript, analysis = case()
        with self.assertRaises(UIValidationError):
            export_document_bytes(transcript, analysis, None)
        transcript["revision"] += 1
        with self.assertRaises(UIValidationError):
            export_document_bytes(transcript, analysis, transcript["revision"])

    def test_synthetic_audio_export_does_not_claim_models_were_skipped(self):
        from docx import Document
        transcript, analysis = case("RU_EXPLICIT")
        content = export_document_bytes(transcript, analysis, transcript["revision"], synthetic_audio=True)
        text = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
        self.assertIn("Синтетическая аудиозапись — выполнена локальная обработка моделями", text)
        self.assertNotIn("модели не запускались", text)


@unittest.skipUnless(importlib.util.find_spec("streamlit"), "Streamlit unavailable; run with requirements-app installed")
class StreamlitWorkflowTests(unittest.TestCase):
    """Real Streamlit widget/state execution; no browser or model quality claim."""

    def start(self, identity=None):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=30).run()
        self.assertEqual(len(app.exception), 0)
        if identity is not None:
            app.radio[0].set_value("Синтетические примеры").run()
            app.checkbox(key="consent").set_value(True).run()
            app.selectbox[0].set_value(identity).run()
            self.button(app, "Открыть вымышленный пример").click().run()
            self.assertEqual(len(app.exception), 0)
        return app

    @staticmethod
    def button(app, label):
        return next(button for button in app.button if button.label == label)

    @staticmethod
    def checkbox(app, prefix):
        return next(item for item in app.checkbox if item.key and item.key.startswith(prefix))

    def test_consent_required_to_load_synthetic_example(self):
        app = self.start()
        app.radio[0].set_value("Синтетические примеры").run()
        self.assertTrue(self.button(app, "Открыть вымышленный пример").disabled)
        self.assertNotIn("transcript", app.session_state)

    def test_full_edit_review_export_workflow(self):
        app = self.start("NO_DUE")
        self.assertTrue(self.button(app, "Подготовить DOCX").disabled)
        self.checkbox(app, "export_review_").set_value(True).run()
        self.button(app, "Подготовить DOCX").click().run()
        self.assertGreater(len(app.session_state["docx_bytes"]), 1000)
        next(item for item in app.text_input if item.label == "Дата ГГГГ-ММ-ДД или пусто").set_value("2026-10-01")
        next(item for item in app.text_input if item.label == "Объяснение исправлений").set_value("Уточнили после встречи")
        self.button(app, "Сохранить правки и сбросить подтверждение").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertNotIn("docx_bytes", app.session_state)
        self.assertNotIn("reviewed_revision", app.session_state)
        self.assertEqual(app.session_state["transcript"]["revision"], 2)
        self.assertEqual(app.session_state["analysis"]["data"]["source_revision"], 1)
        self.assertFalse(any(button.label == "Подготовить DOCX" for button in app.button))
        self.checkbox(app, "content_").set_value(True).run()
        self.button(app, "Завершить проверку исправленного снимка").click().run()
        self.assertFalse(self.checkbox(app, "export_review_").value)
        self.assertTrue(self.button(app, "Подготовить DOCX").disabled)
        self.checkbox(app, "export_review_").set_value(True).run()
        self.button(app, "Подготовить DOCX").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertGreater(len(app.session_state["docx_bytes"]), 1000)
        self.assertIsNone(app.session_state["analysis"]["data"]["tasks"][0]["due_text"])

    def test_mode_switch_discards_synthetic_snapshot(self):
        app = self.start("RU_EXPLICIT")
        self.assertIn("transcript", app.session_state)
        app.radio[0].set_value("Аудиозапись").run()
        self.assertNotIn("transcript", app.session_state)
        self.assertNotIn("analysis", app.session_state)
        self.assertEqual(app.session_state["snapshot_history"][0]["transcript"], case("RU_EXPLICIT")[0])

    def test_transcript_edit_preserves_previous_analysis_and_user_edits(self):
        transcript, current = case("USER_CLARIFIED_DUE")
        app = self.start()
        app.session_state["transcript"] = transcript
        app.session_state["analysis"] = current
        app.session_state["synthetic"] = False
        app.session_state["synthetic_audio"] = True
        app.run()
        segment = next(item for item in app.text_area if item.label.startswith("s1 ·"))
        segment.set_value(segment.value + " Проверка.")
        self.button(app, "Сохранить подтверждения и текст").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertIsNone(app.session_state["analysis"])
        saved = app.session_state["snapshot_history"][0]
        self.assertEqual(saved["transcript"], transcript)
        self.assertEqual(saved["analysis"], current)
        self.assertEqual(len(saved["analysis"]["data"]["user_edits"]), 1)

    def test_metadata_change_preserves_history_and_deletion_requires_confirmation(self):
        transcript, current = case("USER_CLARIFIED_DUE")
        app = self.start()
        app.session_state["transcript"] = transcript
        app.session_state["analysis"] = current
        app.session_state["synthetic"] = False
        app.run()
        next(item for item in app.text_area if item.label.startswith("Участники:")).set_value("Ирина\nДанияр").run()
        self.assertEqual(len(app.exception), 0)
        self.assertNotIn("transcript", app.session_state)
        self.assertEqual(app.session_state["snapshot_history"][0]["analysis"], current)
        self.assertTrue(self.button(app, "Удалить сохранённые версии").disabled)
        self.assertTrue(any(item.label == "Скачать сохранённый снимок JSON" for item in app.get("download_button")))
        self.checkbox(app, "erase_history_").set_value(True).run()
        self.button(app, "Удалить сохранённые версии").click().run()
        self.assertEqual(app.session_state["snapshot_history"], [])

    def test_case_change_discards_previous_result(self):
        app = self.start("RU_EXPLICIT")
        app.selectbox[0].set_value("NO_DUE").run()
        self.assertNotIn("transcript", app.session_state)

    def test_silence_does_not_offer_export(self):
        app = self.start("NO_SPEECH")
        self.assertFalse(any(button.label == "Подготовить DOCX" for button in app.button))

    def test_sample_audio_calls_pipeline_without_expected_result_substitution(self):
        fixture_path = Path(__file__).resolve().parents[1] / "fixtures" / "audio" / "synthetic_ru_two_speakers.wav"
        if not fixture_path.exists():
            self.skipTest("Optional local synthetic WAV unavailable")
        transcript, _ = case("RU_EXPLICIT")
        transcript["segments"][0]["text"] = "Тестовый результат вызванного аудиомодуля."
        app = self.start()
        app.checkbox(key="consent").set_value(True).run()
        self.assertTrue(self.button(app, "Испытать на синтетической аудиозаписи").disabled)
        app.checkbox(key="sample_audio_confirmed").set_value(True).run()
        with patch("audio.pipeline.transcribe", return_value={"status": "ok", "data": transcript, "error": None}) as transcribe:
            self.button(app, "Испытать на синтетической аудиозаписи").click().run()
        self.assertEqual(len(app.exception), 0)
        transcribe.assert_called_once()
        self.assertEqual(Path(transcribe.call_args.args[0]), fixture_path)
        self.assertEqual(transcribe.call_args.kwargs["num_speakers"], 2)
        self.assertIsNone(app.session_state["analysis"])
        self.assertFalse(app.session_state["synthetic"])
        self.assertTrue(app.session_state["synthetic_audio"])
        self.assertEqual(app.session_state["transcript"]["segments"][0]["text"], "Тестовый результат вызванного аудиомодуля.")

    def test_reanalysis_requires_comparison_before_replacing_manual_edits(self):
        transcript, current = case("USER_CLARIFIED_DUE")
        proposed = copy.deepcopy(current)
        proposed["data"]["tasks"][0] = copy.deepcopy(current["data"]["user_edits"][0]["original"])
        proposed["data"]["user_edits"] = []
        app = self.start()
        app.checkbox(key="consent").set_value(True).run()
        app.session_state["transcript"] = transcript
        app.session_state["analysis"] = current
        app.session_state["synthetic"] = False
        app.session_state["synthetic_audio"] = True
        app.run()
        self.checkbox(app, "replace_").set_value(True).run()
        with patch("analysis.local_model.analyze", return_value=proposed):
            self.button(app, "Выделить поручения и подготовить саммари").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["analysis"], current)
        self.assertEqual(app.session_state["pending_analysis"], proposed)
        self.assertTrue(self.button(app, "Принять новый анализ").disabled)
        self.checkbox(app, "accept_new_").set_value(True).run()
        self.button(app, "Принять новый анализ").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["analysis"], proposed)
        self.assertFalse(self.checkbox(app, "export_review_").value)


if __name__ == "__main__":
    unittest.main()
