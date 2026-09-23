"""Real DOCX content and failure checks on invented data in temporary folders.

Reading XML/python-docx here verifies content, not rendered layout or an audio
pipeline. Canonical fixtures and all pre-existing test files remain unchanged.
"""

import contextlib
import copy
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from docx import Document
from docx.document import Document as DocumentClass

import contracts
from export_docx import export_docx


FIXTURE_PATH = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_cases.json"
with FIXTURE_PATH.open(encoding="utf-8") as fixture_file:
    CASES = {case["id"]: case for case in json.load(fixture_file)["cases"]}

try:
    ZoneInfo("Asia/Qyzylorda")
    HAS_TIMEZONE_DATA = True
except ZoneInfoNotFoundError:
    HAS_TIMEZONE_DATA = False


def case_copy(case_id="RU_EXPLICIT", *, unknown_meeting=True):
    case = copy.deepcopy(CASES[case_id])
    if unknown_meeting:
        case["input"]["meeting"] = {"started_at": None, "timezone": None}
    return case["input"], case["expected"]


def document_text(document):
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(texts)


def stress_snapshot():
    """Long synthetic content; deliberately not evidence of model output."""
    transcript, analysis = case_copy()
    alphabet = "Ә Ғ Қ Ң Ө Ұ Ү Һ І ә ғ қ ң ө ұ ү һ і"
    long_text = ("Синтетический фрагмент для проверки переноса длинного текста. " * 8)
    transcript["segments"] = []
    analysis["data"].update(summary=[], decisions=[], tasks=[])
    for index in range(18):
        segment_id = "stress-s" + str(index + 1)
        speaker_id = "spk1" if index % 2 == 0 else "spk2"
        transcript["segments"].append({
            "id": segment_id, "start": index * 10, "end": index * 10 + 8,
            "speaker_id": speaker_id,
            "text": f"Синтетическая реплика {index + 1:02d}. {alphabet}\n{long_text}\tДо проверки — до 30 сентября 2026 года.",
        })
        analysis["data"]["summary"].append({"id": "stress-f" + str(index + 1),
            "text": f"Синтетический пункт саммари {index + 1:02d}. {alphabet}",
            "source_segment_ids": [segment_id]})
        if index % 2 == 0:
            analysis["data"]["tasks"].append({
                "id": "stress-t" + str(index + 1),
                "description": f"Синтетическое поручение {index + 1:02d}. {long_text}",
                "author_speaker_id": speaker_id, "assignee_id": "p2",
                "due_text": "до 30 сентября 2026 года", "due_date": "2026-09-30",
                "source_segment_ids": [segment_id], "requires_clarification": False,
                "clarification_reasons": [],
            })
    return transcript, analysis


class ExportDocxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="qonimai-export-tests-")
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.target = self.directory / "synthetic.docx"

    def assert_error(self, result, code):
        self.assertIsInstance(result, dict)
        self.assertEqual(set(result), {"status", "data", "error"})
        self.assertEqual(result["status"], "error")
        self.assertIsNone(result["data"])
        self.assertEqual(set(result["error"]), {"code", "message", "retryable"})
        self.assertEqual(result["error"]["code"], code)
        self.assertIs(result["error"]["retryable"], False)
        self.assertIsInstance(result["error"]["message"], str)
        self.assertTrue(result["error"]["message"].strip())
        return result

    def export_and_read(self, transcript, analysis, target=None):
        target = self.target if target is None else target
        result = export_docx(transcript, analysis, transcript["revision"], target)
        self.assertEqual(set(result), {"status", "data", "error"})
        self.assertEqual(result["status"], "ok", result)
        self.assertIsNone(result["error"])
        self.assertEqual(set(result["data"]), {"path", "source_revision"})
        self.assertEqual(result["data"]["source_revision"], transcript["revision"])
        self.assertIsInstance(result["data"]["path"], str)
        ready_path = Path(result["data"]["path"])
        self.assertTrue(ready_path.is_absolute())
        self.assertEqual(ready_path.resolve(), Path(target).resolve())
        self.assertTrue(ready_path.is_file())
        return Document(ready_path), result

    def assert_no_write(self, transcript, analysis, reviewed, code):
        for existing in (False, True):
            with self.subTest(existing=existing):
                if self.target.exists():
                    self.target.unlink()
                original = b"SYNTHETIC_PREVIOUS_DOCUMENT_CONTENT"
                if existing:
                    self.target.write_bytes(original)
                before = set(self.directory.iterdir())
                self.assert_error(export_docx(transcript, analysis, reviewed, self.target), code)
                self.assertEqual(set(self.directory.iterdir()), before)
                if existing:
                    self.assertEqual(self.target.read_bytes(), original)
                else:
                    self.assertFalse(self.target.exists())

    def test_success_returns_readable_docx_and_snapshot_revision(self):
        transcript, analysis = case_copy()
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Протокол совещания", text)
        self.assertIn("Саммари", text)
        self.assertIn("Принятые решения", text)
        self.assertIn("Поручения", text)
        self.assertIn("Транскрипт", text)

    def test_string_and_pathlike_targets_are_supported(self):
        transcript, analysis = case_copy()
        self.export_and_read(transcript, analysis, str(self.target))
        self.export_and_read(transcript, analysis, self.directory / "pathlike.docx")

    def test_current_transcript_names_and_order_are_preserved(self):
        transcript, analysis = case_copy()
        transcript["participants"][0]["name"] = "Мария"
        transcript["speakers"][0]["display_name"] = "Мария"
        transcript["segments"][1]["text"] = "Актуальный текст после ручной проверки.\nВторая строка.\tТабуляция."
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Мария", text)
        positions = []
        for segment in transcript["segments"]:
            self.assertIn(segment["id"], text)
            self.assertIn(segment["text"], text)
            positions.append(text.index(segment["text"]))
        self.assertEqual(positions, sorted(positions))
        self.assertNotIn("Да, подготовлю таблицу.", text)

    @unittest.skipUnless(HAS_TIMEZONE_DATA, "Real IANA timezone database unavailable")
    def test_meeting_metadata_comes_from_snapshot(self):
        transcript, analysis = case_copy(unknown_meeting=False)
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("2026-09-23", text)
        self.assertIn("Asia/Qyzylorda", text)
        self.assertIn("2026-09-30", text)

    def test_unknown_meeting_is_explicit_not_system_date(self):
        transcript, analysis = case_copy("MISSING_CONTEXT")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertRegex(text, r"Дата совещания[^\n]*не указан[ао]?")
        self.assertRegex(text, r"Часовой пояс[^\n]*не указан[ао]?")
        self.assertIn(analysis["data"]["tasks"][0]["clarification_reasons"][0]["message"], text)

    def test_task_fields_show_author_and_independent_assignee(self):
        transcript, analysis = case_copy()
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertRegex(text, r"Исполнитель[^\n]*Данияр")
        self.assertRegex(text, r"Автор поручения[^\n]*Ирина")
        task = analysis["data"]["tasks"][0]
        self.assertIn(task["description"], text)
        self.assertIn(task["due_text"], text)
        self.assertIn(task["due_date"], text)
        self.assertRegex(text, r"Источники[^\n]*s1[^\n]*s2")

    def test_fractional_timecodes_and_overlap_are_preserved(self):
        transcript, analysis = case_copy()
        transcript["segments"][0].update(start=1.25, end=5.5)
        transcript["segments"][1].update(start=4.75, end=8.125)
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertRegex(text, r"s1[^\n]*1\.25[^\n]*5\.5")
        self.assertRegex(text, r"s2[^\n]*4\.75[^\n]*8\.125")

    def test_unknown_due_exports_with_visible_clarification(self):
        transcript, analysis = case_copy("NO_DUE")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Требует уточнения", text)
        self.assertIn(analysis["data"]["tasks"][0]["clarification_reasons"][0]["message"], text)
        self.assertRegex(text, r"Исходный срок[^\n]*не указан")
        self.assertRegex(text, r"Нормализованная дата[^\n]*не указан[ао]?")
        self.assertNotIn("null", text)

    def test_unknown_assignee_is_not_replaced_with_author(self):
        transcript, analysis = case_copy("UNKNOWN_ASSIGNEE")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertRegex(text, r"Исполнитель[^\n]*не определ[её]н")
        self.assertIn("Требует уточнения", text)
        self.assertIn(analysis["data"]["tasks"][0]["clarification_reasons"][0]["message"], text)
        self.assertNotRegex(text, r"Исполнитель[^\n]*Ирина")

    def test_no_tasks_is_stated_explicitly(self):
        transcript, analysis = case_copy("NO_TASKS")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Поручения не выявлены", text)
        self.assertIn(analysis["data"]["summary"][0]["text"], text)

    def test_unconfirmed_speaker_hint_is_not_a_confirmed_identity(self):
        transcript, analysis = case_copy("UNCONFIRMED_SELF")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("spk2", text)
        self.assertRegex(text.lower(), r"неподтвержд|не подтвержд")
        self.assertRegex(text, r"Исполнитель[^\n]*не определ[её]н")
        participants_section = text.split("Подтверждённые участники", 1)[1].split("Саммари", 1)[0]
        self.assertNotIn("Данияр", participants_section)

    def test_user_clarified_due_preserves_missing_original(self):
        transcript, analysis = case_copy("USER_CLARIFIED_DUE")
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Уточнено пользователем", text)
        self.assertIn(analysis["data"]["user_edits"][0]["reason"], text)
        self.assertIn("2026-09-30", text)
        self.assertRegex(text, r"Исходный срок[^\n]*не указан")
        self.assertIn("Исходное значение", text)
        self.assertIn("Актуальное значение", text)
        self.assertNotIn("null", text)

    def test_all_user_edit_types_keep_original_current_and_reason(self):
        transcript, analysis = case_copy()
        transcript["revision"] = 2
        analysis["data"]["source_revision"] = 2
        analysis["data"]["decisions"] = [{"id": "d1", "text": "Подготовить таблицу расходов.", "source_segment_ids": ["s1"]}]
        edits = []
        for target_type, collection, field, current_text in (
            ("summary", "summary", "text", "Уточнённое пользователем саммари о таблице."),
            ("decision", "decisions", "text", "Уточнённое пользователем решение о таблице."),
            ("task", "tasks", "description", "Уточнить перечень расходов в таблице"),
        ):
            item = analysis["data"][collection][0]
            original = copy.deepcopy(item)
            item[field] = current_text
            fields = [field]
            if target_type == "task":
                item.update(assignee_id="p1", due_date="2026-10-02")
                fields.extend(["assignee_id", "due_date"])
            edits.append({"target_type": target_type, "target_id": item["id"], "original": original,
                          "changed_fields": fields, "reason": "Синтетическое объяснение: " + target_type})
        analysis["data"]["user_edits"] = edits
        self.assertEqual(contracts.validate_export(transcript, analysis, 2), [])
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertGreaterEqual(text.count("Уточнено пользователем"), 3)
        for edit in edits:
            self.assertIn(edit["reason"], text)
            field = "description" if edit["target_type"] == "task" else "text"
            self.assertIn(edit["original"][field], text)
        self.assertIn("Уточнённое пользователем саммари о таблице.", text)
        self.assertIn("Уточнённое пользователем решение о таблице.", text)
        self.assertIn("Уточнить перечень расходов в таблице", text)
        self.assertIn("2026-10-02", text)
        self.assertIn("2026-09-30", text)

    def test_long_synthetic_content_preserves_unicode_newlines_and_task_order(self):
        transcript, analysis = stress_snapshot()
        self.assertEqual(contracts.validate_export(transcript, analysis, 1), [])
        document, _ = self.export_and_read(transcript, analysis)
        text = document_text(document)
        self.assertIn("Ә Ғ Қ Ң Ө Ұ Ү Һ І ә ғ қ ң ө ұ ү һ і", text)
        positions = []
        for segment in transcript["segments"]:
            self.assertIn(segment["text"], text)
            positions.append(text.index(segment["text"]))
        self.assertEqual(positions, sorted(positions))
        positions = [text.index(task["description"]) for task in analysis["data"]["tasks"]]
        self.assertEqual(positions, sorted(positions))
        self.assertGreater(len(document.paragraphs), 30)

    def test_body_font_and_builtin_headings_are_readable(self):
        transcript, analysis = case_copy()
        document, _ = self.export_and_read(transcript, analysis)
        self.assertIsNotNone(document.styles["Normal"].font.size)
        self.assertGreaterEqual(document.styles["Normal"].font.size.pt, 11)
        self.assertLessEqual(document.styles["Normal"].font.size.pt, 12)
        headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertTrue(headings)
        for paragraph in headings:
            self.assertEqual(str(paragraph.style.font.color.rgb), "000000")

    def test_missing_review_never_writes_or_corrupts_existing_target(self):
        transcript, analysis = case_copy()
        self.assert_no_write(transcript, analysis, None, "REVIEW_REQUIRED")

    def test_stale_analysis_never_writes_or_corrupts_existing_target(self):
        transcript, analysis = case_copy()
        transcript["revision"] = 2
        self.assert_no_write(transcript, analysis, 2, "STALE_ANALYSIS")

    def test_broken_reference_never_writes_or_corrupts_existing_target(self):
        transcript, analysis = case_copy()
        analysis["data"]["tasks"][0]["source_segment_ids"] = ["missing"]
        self.assert_no_write(transcript, analysis, 1, "INVALID_MODEL_OUTPUT")

    def test_model_error_never_writes_or_corrupts_existing_target(self):
        transcript, _ = case_copy()
        analysis = {"status": "error", "data": None,
                    "error": {"code": "MODEL_UNAVAILABLE", "message": "Синтетическая ошибка", "retryable": False}}
        self.assert_no_write(transcript, analysis, 1, "INVALID_MODEL_OUTPUT")

    def test_no_speech_never_writes_even_after_review(self):
        transcript, analysis = case_copy("NO_SPEECH")
        self.assert_no_write(transcript, analysis, 1, "NO_CONTENT")

    def test_invalid_transcript_never_writes_or_corrupts_existing_target(self):
        transcript, analysis = case_copy()
        transcript["revision"] = True
        self.assert_no_write(transcript, analysis, 1, "INVALID_INPUT")

    def test_validation_precedes_any_target_path_processing(self):
        class UnusablePath(os.PathLike):
            def __fspath__(self):
                raise AssertionError("Target path must not be inspected before validation")
        transcript, analysis = case_copy()
        self.assert_error(export_docx(transcript, analysis, None, UnusablePath()), "REVIEW_REQUIRED")
        self.assertEqual(list(self.directory.iterdir()), [])

    def test_error_precedence_is_deterministic(self):
        transcript, analysis = case_copy()
        transcript["revision"] = 2
        self.assert_error(export_docx(transcript, analysis, None, self.target), "STALE_ANALYSIS")
        analysis["data"]["tasks"][0]["assignee_id"] = "missing"
        self.assert_error(export_docx(transcript, analysis, None, self.target), "INVALID_MODEL_OUTPUT")
        transcript["revision"] = False
        self.assert_error(export_docx(transcript, analysis, None, self.target), "INVALID_INPUT")
        transcript, analysis = case_copy("NO_SPEECH")
        self.assert_error(export_docx(transcript, analysis, None, self.target), "NO_CONTENT")
        self.assertFalse(self.target.exists())

    def test_simulated_missing_iana_database_is_not_export_failure(self):
        transcript, analysis = case_copy(unknown_meeting=False)
        with mock.patch("contracts.ZoneInfo", side_effect=ZoneInfoNotFoundError("Synthetic database failure")):
            self.assert_no_write(transcript, analysis, 1, "TIMEZONE_DATA_UNAVAILABLE")

    def test_invalid_path_types_and_extensions_return_export_failed(self):
        transcript, analysis = case_copy()
        for target in (None, False, 123, [], {}, "", self.directory / "wrong.txt", self.directory / "no-extension"):
            with self.subTest(target=str(target)):
                self.assert_error(export_docx(transcript, analysis, 1, target), "EXPORT_FAILED")
                self.assertEqual(list(self.directory.iterdir()), [])

    def test_urls_and_unc_paths_are_rejected_without_writes(self):
        transcript, analysis = case_copy()
        paths = ["https://example.invalid/synthetic.docx", "file:///tmp/synthetic.docx",
                 "\\\\synthetic-server\\share\\synthetic.docx", "//synthetic-server/share/synthetic.docx"]
        for target in paths:
            with self.subTest(target=target):
                self.assert_error(export_docx(transcript, analysis, 1, target), "EXPORT_FAILED")
                self.assertEqual(list(self.directory.iterdir()), [])

    def test_missing_parent_is_not_created(self):
        transcript, analysis = case_copy()
        target = self.directory / "uncreated" / "synthetic.docx"
        self.assert_error(export_docx(transcript, analysis, 1, target), "EXPORT_FAILED")
        self.assertFalse(target.parent.exists())

    def test_directory_target_is_not_replaced(self):
        transcript, analysis = case_copy()
        self.target.mkdir()
        sentinel = self.target / "keep.txt"
        sentinel.write_text("Сохранить", encoding="utf-8")
        self.assert_error(export_docx(transcript, analysis, 1, self.target), "EXPORT_FAILED")
        self.assertEqual(sentinel.read_text(encoding="utf-8"), "Сохранить")

    def test_success_atomically_replaces_existing_target_from_same_directory(self):
        transcript, analysis = case_copy()
        original = b"SYNTHETIC_EXISTING_DOCUMENT"
        self.target.write_bytes(original)
        real_replace = os.replace
        replaced_paths = []
        def inspect_replace(source, destination):
            source, destination = Path(source), Path(destination)
            self.assertEqual(source.parent.resolve(), self.directory.resolve())
            self.assertEqual(destination.resolve(), self.target.resolve())
            self.assertEqual(self.target.read_bytes(), original)
            self.assertIn("Протокол совещания", document_text(Document(source)))
            replaced_paths.append(source)
            return real_replace(source, destination)
        with mock.patch("export_docx.os.replace", side_effect=inspect_replace):
            self.export_and_read(transcript, analysis)
        self.assertEqual(len(replaced_paths), 1)
        self.assertFalse(replaced_paths[0].exists())
        self.assertEqual(set(self.directory.iterdir()), {self.target})

    def test_save_failure_cleans_only_own_temporary_file(self):
        transcript, analysis = case_copy()
        original = b"SYNTHETIC_EXISTING_DOCUMENT"
        self.target.write_bytes(original)
        unrelated = self.directory / "qonimai-unrelated.tmp"
        unrelated.write_bytes(b"KEEP")
        def fail_after_partial_save(_document, destination):
            if hasattr(destination, "write"):
                destination.write(b"PARTIAL SYNTHETIC DOCX")
            else:
                Path(destination).write_bytes(b"PARTIAL SYNTHETIC DOCX")
            raise OSError("SYNTHETIC_PRIVATE_SAVE_ERROR")
        with mock.patch("docx.document.Document.save", autospec=True, side_effect=fail_after_partial_save):
            self.assert_error(export_docx(transcript, analysis, 1, self.target), "EXPORT_FAILED")
        self.assertEqual(self.target.read_bytes(), original)
        self.assertEqual(unrelated.read_bytes(), b"KEEP")
        self.assertEqual(set(self.directory.iterdir()), {self.target, unrelated})

    def test_replace_failure_preserves_existing_target_and_cleans_temp(self):
        transcript, analysis = case_copy()
        self.target.write_bytes(b"KEEP ORIGINAL")
        with mock.patch("export_docx.os.replace", side_effect=OSError("SYNTHETIC_PRIVATE_REPLACE_ERROR")):
            self.assert_error(export_docx(transcript, analysis, 1, self.target), "EXPORT_FAILED")
        self.assertEqual(self.target.read_bytes(), b"KEEP ORIGINAL")
        self.assertEqual(set(self.directory.iterdir()), {self.target})

    def test_failed_new_export_leaves_no_partial_target(self):
        transcript, analysis = case_copy()
        with mock.patch("export_docx.os.replace", side_effect=OSError("Synthetic replace error")):
            self.assert_error(export_docx(transcript, analysis, 1, self.target), "EXPORT_FAILED")
        self.assertEqual(list(self.directory.iterdir()), [])

    def test_corrupt_save_without_exception_is_not_published(self):
        transcript, analysis = case_copy()
        self.target.write_bytes(b"KEEP ORIGINAL")
        def silently_corrupt_save(_document, destination):
            if hasattr(destination, "write"):
                destination.write(b"BROKEN SYNTHETIC DOCX")
            else:
                Path(destination).write_bytes(b"BROKEN SYNTHETIC DOCX")
        with mock.patch("docx.document.Document.save", autospec=True, side_effect=silently_corrupt_save):
            self.assert_error(export_docx(transcript, analysis, 1, self.target), "EXPORT_FAILED")
        self.assertEqual(self.target.read_bytes(), b"KEEP ORIGINAL")
        self.assertEqual(set(self.directory.iterdir()), {self.target})

    def test_process_interrupts_propagate_and_cleanup_temporary_file(self):
        transcript, analysis = case_copy()
        self.target.write_bytes(b"KEEP ORIGINAL")
        for exception in (KeyboardInterrupt, SystemExit):
            with self.subTest(exception=exception.__name__):
                with mock.patch("docx.document.Document.save", side_effect=exception):
                    with self.assertRaises(exception):
                        export_docx(transcript, analysis, 1, self.target)
                self.assertEqual(self.target.read_bytes(), b"KEEP ORIGINAL")
                self.assertEqual(set(self.directory.iterdir()), {self.target})

    def test_inputs_remain_unchanged_after_success_and_failure(self):
        transcript, analysis = case_copy("USER_CLARIFIED_DUE")
        before = copy.deepcopy((transcript, analysis))
        self.export_and_read(transcript, analysis)
        self.assertEqual((transcript, analysis), before)
        self.assert_error(export_docx(transcript, analysis, None, self.target), "REVIEW_REQUIRED")
        self.assertEqual((transcript, analysis), before)
        with mock.patch("export_docx.os.replace", side_effect=OSError("Synthetic failure")):
            self.assert_error(export_docx(transcript, analysis, 2, self.target), "EXPORT_FAILED")
        self.assertEqual((transcript, analysis), before)

    def test_errors_and_logs_do_not_expose_private_text_path_or_exception(self):
        marker = "SYNTHETIC_PRIVATE_MARKER"
        transcript, analysis = case_copy()
        transcript["segments"][1]["text"] = marker
        target = self.directory / (marker + ".docx")
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr), mock.patch("logging.Logger._log") as logger:
            with mock.patch("docx.document.Document.save", side_effect=OSError(marker)):
                result = export_docx(transcript, analysis, 1, target)
        self.assert_error(result, "EXPORT_FAILED")
        serialized = json.dumps(result, ensure_ascii=False)
        self.assertNotIn(marker, serialized)
        self.assertNotIn(str(self.directory), serialized)
        self.assertEqual(stdout.getvalue(), "")
        self.assertEqual(stderr.getvalue(), "")
        logger.assert_not_called()
        self.assertEqual(list(self.directory.iterdir()), [])

    def test_cli_requires_explicit_review_and_keeps_existing_file(self):
        from examples.export_fixture import main, SYNTHETIC_NOTICE
        self.target.write_bytes(b"KEEP ORIGINAL")
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            code = main(["--case", "USER_CLARIFIED_DUE", "--output", str(self.target)])
        self.assertNotEqual(code, 0)
        self.assertIn(SYNTHETIC_NOTICE, stdout.getvalue())
        self.assertIn("--confirm-reviewed", stderr.getvalue())
        self.assertEqual(self.target.read_bytes(), b"KEEP ORIGINAL")
        self.assertEqual(set(self.directory.iterdir()), {self.target})

    @unittest.skipUnless(HAS_TIMEZONE_DATA, "Real IANA timezone database unavailable")
    def test_cli_creates_labelled_canonical_and_stress_documents(self):
        from examples.export_fixture import main, SYNTHETIC_NOTICE, build_stress_case
        for case_id in ("USER_CLARIFIED_DUE", "STRESS"):
            with self.subTest(case_id=case_id):
                target = self.directory / ("synthetic_" + case_id.lower() + ".docx")
                stdout, stderr = io.StringIO(), io.StringIO()
                with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                    code = main(["--case", case_id, "--output", str(target), "--confirm-reviewed"])
                self.assertEqual(code, 0, stderr.getvalue())
                self.assertIn(SYNTHETIC_NOTICE, stdout.getvalue())
                document = Document(target)
                text = document_text(document)
                self.assertIn(SYNTHETIC_NOTICE, text)
                if case_id == "STRESS":
                    case = build_stress_case()
                    self.assertIn("Ә Ғ Қ Ң Ө Ұ Ү Һ І ә ғ қ ң ө ұ ү һ і", text)
                    for segment in case["input"]["segments"]:
                        self.assertIn(segment["text"], text)
                    for edit in case["expected"]["data"]["user_edits"]:
                        self.assertIn(edit["reason"], text)

    @unittest.skipUnless(HAS_TIMEZONE_DATA, "Real IANA timezone database unavailable")
    def test_cli_no_speech_still_cannot_export_after_review(self):
        from examples.export_fixture import main
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            code = main(["--case", "NO_SPEECH", "--output", str(self.target), "--confirm-reviewed"])
        self.assertNotEqual(code, 0)
        self.assertEqual(list(self.directory.iterdir()), [])

    @unittest.skipUnless(HAS_TIMEZONE_DATA, "Real IANA timezone database unavailable")
    def test_cli_label_save_failure_preserves_existing_target(self):
        from examples.export_fixture import main
        self.target.write_bytes(b"KEEP ORIGINAL")
        save_calls = 0
        real_save = DocumentClass.save
        def fail_second_save(document, destination):
            nonlocal save_calls
            save_calls += 1
            if save_calls == 2:
                raise OSError("SYNTHETIC_PRIVATE_CLI_SAVE_ERROR")
            return real_save(document, destination)
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            with mock.patch("docx.document.Document.save", autospec=True, side_effect=fail_second_save):
                code = main(["--case", "USER_CLARIFIED_DUE", "--output", str(self.target), "--confirm-reviewed"])
        self.assertNotEqual(code, 0)
        self.assertEqual(save_calls, 2)
        self.assertEqual(self.target.read_bytes(), b"KEEP ORIGINAL")
        self.assertEqual(set(self.directory.iterdir()), {self.target})
        self.assertNotIn("SYNTHETIC_PRIVATE_CLI_SAVE_ERROR", stdout.getvalue() + stderr.getvalue())


def canonical_test(case_id):
    def test(self):
        transcript, analysis = case_copy(case_id, unknown_meeting=False)
        if transcript["meeting"]["timezone"] is not None and not HAS_TIMEZONE_DATA:
            self.skipTest("Real IANA timezone database unavailable; canonical fixture unchanged")
        before = copy.deepcopy((transcript, analysis))
        if case_id == "NO_SPEECH":
            self.assert_no_write(transcript, analysis, transcript["revision"], "NO_CONTENT")
        else:
            document, _ = self.export_and_read(transcript, analysis)
            text = document_text(document)
            for participant in transcript["participants"]:
                self.assertIn(participant["name"], text)
            for segment in transcript["segments"]:
                self.assertIn(segment["text"], text)
                self.assertIn(segment["id"], text)
            for collection in ("summary", "decisions"):
                for item in analysis["data"][collection]:
                    self.assertIn(item["text"], text)
            for task in analysis["data"]["tasks"]:
                self.assertIn(task["description"], text)
                for field in ("due_text", "due_date"):
                    if task[field] is not None:
                        self.assertIn(task[field], text)
                for reason in task["clarification_reasons"]:
                    self.assertIn(reason["message"], text)
        self.assertEqual((transcript, analysis), before)
    test.__doc__ = "Actual DOCX export of canonical synthetic snapshot: " + case_id
    return test


for case_id in CASES:
    setattr(ExportDocxTests, "test_canonical_" + case_id.lower(), canonical_test(case_id))


if __name__ == "__main__":
    unittest.main()
